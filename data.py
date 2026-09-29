"""Local multi-format ingestion. No uploaded code is executed."""
import io
import json
import re
import sqlite3
import tempfile
import zipfile
from pathlib import Path

import pandas as pd

EXTENSIONS = ['csv', 'tsv', 'xlsx', 'xls', 'ods', 'json', 'jsonl', 'ndjson',
              'parquet', 'feather', 'sqlite', 'sqlite3', 'db', 'pdf', 'docx',
              'txt', 'md', 'xml']
MAX_ROWS = 100_000
MAX_TABLES = 30
MAX_BYTES = 20 * 1024 * 1024


def identifier(value):
    name = re.sub(r'[^a-z0-9_]', '_', str(value).lower()).strip('_')[:55]
    return ('t_' + name if name[:1].isdigit() else name) or 'unnamed'


def unique_name(value, used):
    base = identifier(value)
    name, suffix = base, 2
    while name in used:
        name = f'{base}_{suffix}'
        suffix += 1
    used.add(name)
    return name


def clean_frame(frame):
    if len(frame) > MAX_ROWS or len(frame.columns) > 150:
        raise ValueError('Limit: 100,000 rows and 150 columns per table. Split the file first.')
    if not len(frame.columns):
        raise ValueError('No columns found.')
    frame = frame.copy()
    used = set()
    frame.columns = [unique_name(c, used) for c in frame.columns]
    for col in frame:
        if pd.api.types.is_datetime64_any_dtype(frame[col]):
            frame[col] = frame[col].dt.strftime('%Y-%m-%d %H:%M:%S')
        elif not pd.api.types.is_numeric_dtype(frame[col]):
            frame[col] = frame[col].map(
                lambda v: json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list))
                else v.hex() if isinstance(v, bytes) else v
            )
    return frame


def decode(data):
    for encoding in ('utf-8-sig', 'utf-16', 'cp1252'):
        try:
            return data.decode(encoding)
        except UnicodeError:
            pass
    raise ValueError('Unsupported text encoding. Save the file as UTF-8.')


def check_archive(data):
    if zipfile.is_zipfile(io.BytesIO(data)):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Expanded archive exceeds 100 MB. Use a smaller file.')


def text_rows(text, source, page):
    return [{'source': source, 'page': page, 'chunk': i // 1800 + 1,
             'text': text[i:i + 1800]} for i in range(0, len(text), 1800)
            if text[i:i + 1800].strip()]


def load_file(filename, data):
    """Return named DataFrames and visible warnings; reject oversized inputs."""
    if len(data) > MAX_BYTES:
        raise ValueError('Each file must be 20 MB or smaller.')
    ext = Path(filename).suffix.lower().lstrip('.')
    stem = identifier(Path(filename).stem)
    stream = io.BytesIO(data)
    tables, notes = {}, []
    if ext in ('csv', 'tsv'):
        tables[stem] = pd.read_csv(io.StringIO(decode(data)), sep='\t' if ext == 'tsv' else None,
                                   engine='python', nrows=MAX_ROWS + 1)
    elif ext in ('xlsx', 'xls', 'ods'):
        check_archive(data)
        with pd.ExcelFile(stream) as book:
            if len(book.sheet_names) > MAX_TABLES:
                raise ValueError('Workbook has more than 30 sheets.')
            for sheet in book.sheet_names:
                tables[f'{stem}_{sheet}'] = pd.read_excel(book, sheet_name=sheet, nrows=MAX_ROWS + 1)
    elif ext in ('jsonl', 'ndjson'):
        tables[stem] = pd.read_json(stream, lines=True)
    elif ext == 'json':
        obj = json.loads(decode(data))
        if isinstance(obj, dict) and obj and all(isinstance(v, list) for v in obj.values()):
            if all(not v or not isinstance(v[0], dict) for v in obj.values()):
                tables[stem] = pd.DataFrame(obj)
            else:
                for key, rows in obj.items():
                    tables[f'{stem}_{key}'] = pd.json_normalize(rows) if rows and isinstance(rows[0], dict) else pd.DataFrame({'value': rows})
        else:
            tables[stem] = pd.json_normalize(obj if isinstance(obj, list) else [obj])
    elif ext in ('parquet', 'feather'):
        tables[stem] = pd.read_parquet(stream) if ext == 'parquet' else pd.read_feather(stream)
    elif ext in ('sqlite', 'sqlite3', 'db'):
        # Copy physical tables only; never execute uploaded triggers/views.
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.db'
            path.write_bytes(data)
            conn = sqlite3.connect(path.as_uri() + '?mode=ro', uri=True)
            try:
                conn.execute('PRAGMA trusted_schema=OFF')
                names = conn.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'").fetchall()
                if len(names) > MAX_TABLES:
                    raise ValueError('Database has more than 30 tables.')
                for (name,) in names:
                    quoted = '"' + name.replace('"', '""') + '"'
                    tables[f'{stem}_{name}'] = pd.read_sql_query(f'SELECT * FROM {quoted} LIMIT {MAX_ROWS + 1}', conn)
            finally:
                conn.close()
    elif ext == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(stream)
        if len(reader.pages) > 200:
            raise ValueError('PDF limit: 200 pages.')
        rows = []
        for n, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            rows.extend(text_rows(text, filename, n))
            if not text.strip():
                notes.append(f'Page {n}: no extractable text; OCR may be required.')
        if not rows:
            raise ValueError('No extractable PDF text. Run OCR or upload CSV/XLSX instead.')
        tables[f'{stem}_documents'] = pd.DataFrame(rows)
        notes.append('PDF loaded as text chunks, not reliable numeric tables. Use CSV/XLSX for exact financial aggregation.')
    elif ext == 'docx':
        from docx import Document
        check_archive(data)
        doc = Document(stream)
        rows = text_rows('\n'.join(p.text for p in doc.paragraphs), filename, 1)
        if rows:
            tables[f'{stem}_documents'] = pd.DataFrame(rows)
        for n, table in enumerate(doc.tables, 1):
            values = [[cell.text for cell in row.cells] for row in table.rows]
            if values:
                tables[f'{stem}_table_{n}'] = pd.DataFrame(values[1:], columns=values[0])
        notes.append('Word table cells remain text; convert numeric columns in Data explorer before aggregating.')
    elif ext in ('txt', 'md'):
        tables[f'{stem}_documents'] = pd.DataFrame(text_rows(decode(data), filename, 1))
    elif ext == 'xml':
        from defusedxml.ElementTree import fromstring
        root = fromstring(data)
        tables[stem] = pd.DataFrame([
            {child.tag: child.text for child in record} | dict(record.attrib)
            for record in root
        ])
        notes.append('XML expects repeated flat records. Nested XML should be converted to JSON or CSV.')
    else:
        raise ValueError('Unsupported format. See the supported formats in the guide.')
    if not tables or len(tables) > MAX_TABLES:
        raise ValueError('No readable tables found, or the file exceeds 30 tables.')
    result, used = {}, set()
    for name, frame in tables.items():
        if not len(frame.columns):
            notes.append(f'Skipped empty sheet/table: {name}')
            continue
        result[unique_name(name, used)] = clean_frame(frame)
    return result, notes


def demo_data():
    import random
    rng = random.Random(42)
    rows = []
    for i in range(900):
        date = pd.Timestamp('2026-01-01') + pd.Timedelta(days=rng.randrange(180))
        category = rng.choice(['Technology', 'Fashion', 'Home', 'Beauty'])
        units = rng.randint(1, 8)
        revenue = round(units * rng.uniform(20, 250), 2)
        rows.append({'order_id': i + 1, 'date': str(date.date()), 'category': category,
                     'region': rng.choice(['North', 'South', 'East', 'West']),
                     'country': rng.choice(['India', 'United States', 'Germany', 'Australia']),
                     'channel': rng.choice(['Online', 'Retail', 'Partner']), 'units': units,
                     'revenue': revenue, 'profit': round(revenue * rng.uniform(.1, .35), 2)})
    return {'sales': pd.DataFrame(rows)}
